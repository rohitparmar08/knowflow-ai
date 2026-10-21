import os
from pypdf import PdfReader
import docx


def extract_text(document) -> str:
    """
    Extracts plain text from an uploaded document (PDF, DOCX, or TXT).
    Raises ValueError if extracted text is empty (e.g. scanned/image-only PDF).
    """
    ext = (document.file_type or "").lower().lstrip(".")
    if not ext and document.file:
        ext = os.path.splitext(document.file.name)[1].lower().lstrip(".")

    document.file.open("rb")
    try:
        if ext == "pdf":
            reader = PdfReader(document.file)
            pages_text = []
            for page in reader.pages:
                t = page.extract_text()
                if t and t.strip():
                    pages_text.append(t.strip())
            text = "\n\n".join(pages_text).strip()

        elif ext == "docx":
            doc = docx.Document(document.file)
            parts = []
            for p in doc.paragraphs:
                if p.text and p.text.strip():
                    parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text and cell.text.strip():
                            parts.append(cell.text.strip())
            text = "\n\n".join(parts).strip()

        elif ext == "txt":
            raw_bytes = document.file.read()
            text = raw_bytes.decode("utf-8", errors="replace").strip()

        else:
            raise ValueError(f"Unsupported file type for extraction: .{ext}")

    finally:
        document.file.close()

    if not text:
        raise ValueError(
            f"Extracted text is empty or document contains no readable text: {document.title}"
        )

    return text
